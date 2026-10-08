"""一条命令跑通一场面试陪练（产物落 results/ 与 runs/）。

用法示例：
  python run_demo.py --jd data/jd_sample.md --resume data/resume_sample.md \
      --answers data/answers_demo.json --questions 3 --env-file <你的 .env 路径>
  python run_demo.py --jd ... --resume ... --interactive        # 真人练习模式
  python run_demo.py --jd ... --resume ... --engine langgraph   # 用 LangGraph 编排
  COACH_OFFLINE=1 python run_demo.py ... --questions 2          # 离线确定性模式（不调模型）
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path

from coach import config
from coach.eval import harness
from coach.graph import CoachEngine


def _read(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="ignore")


def _docx_text(path: str) -> str:
    """简历一般是 .docx，这里直接抽文本（不依赖外部服务）。"""
    import docx  # python-docx
    d = docx.Document(path)
    parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for tb in d.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def load_text(path: str) -> str:
    p = Path(path)
    if p.suffix.lower() == ".docx":
        return _docx_text(str(p))
    return _read(str(p))


def load_answers(spec: str | None) -> tuple[list[str], dict]:
    """返回 (顺序回答列表, 按类别回答映射)。两种输入都支持，便于演示与真人练习。"""
    if not spec:
        return [], {}
    p = Path(spec)
    if not p.exists():
        return [], {}
    txt = p.read_text(encoding="utf-8", errors="ignore")
    if p.suffix.lower() == ".json":
        data = json.loads(txt)
        if isinstance(data, dict):
            if "by_category" in data:
                merged = {str(k): str(v) for k, v in data["by_category"].items()}
                merged.update({str(k): str(v) for k, v in data.get("by_qid", {}).items()})
                return [], merged
            data = data.get("answers", [])
        return [str(x) for x in data], {}
    return [blk.strip() for blk in txt.split("\n\n") if blk.strip()], {}


def check_inputs(jd: str, resume: str, answers: str | None) -> list[str]:
    """输入体检：文件不存在/为空/答不上的情况给出人话提示，而不是让用户吃一屏 traceback。"""
    problems: list[str] = []
    for role, path in (("JD 文件", jd), ("简历文件", resume), ("回答文件", answers)):
        if not path:
            if role == "回答文件":
                continue
            problems.append(f"{role}：未提供（--jd / --resume 是必填）")
            continue
        p = Path(path)
        if not p.exists():
            hint = ""
            if p.suffix.lower() == ".docx":
                hint = "（简历建议用 .docx / .md / .txt；路径里有空格要加引号）"
            problems.append(f"{role}：找不到 {path}{hint}")
        elif p.is_file() and p.stat().st_size == 0:
            problems.append(f"{role}：文件是空的 {path}")
    return problems


def print_problems(problems: list[str]) -> None:
    print("\n" + "=" * 70)
    print("× 没法开跑，先解决下面这几项：")
    for i, p in enumerate(problems, 1):
        print(f"  {i}. {p}")
    print("-" * 70)
    print("可以这样改：")
    print("  · 用示例文件先跑通一遍（确认环境没问题）：")
    print("      --jd data/jd_sample.md --resume data/resume_sample.md")
    print("  · 用自己的简历：把路径用英文双引号包起来，例如")
    print('      --resume "C:\\Users\\hexin\\Desktop\\1简历\\校招简历.docx"')
    print("  · 路径不确定时，在文件资源管理器里右键该文件 → 复制文件地址")
    print("=" * 70 + "\n")


def interactive_provider(q) -> str:
    print("\n" + "=" * 70)
    print(f"[{q.qid} | {q.category} | {q.difficulty}] {q.text}")
    print("自评分参考要点：")
    for dim in q.rubric.dims:
        print(f"  - {dim.name}: {'、'.join(dim.keywords)}")
    print("-" * 70 + "（输入 END 结束本题）")
    lines: list[str] = []
    while True:
        try:
            line = input("> ")
        except EOFError:
            break
        if line.strip() == "END":
            break
        lines.append(line)
    return "\n".join(lines).strip()


def preload_env_file(argv: list[str]) -> tuple[str | None, int]:
    """在解析参数之前先把 --env-file 读进来。

    原因：argparse 的 default= 是即时求值，而 COACH_MAX_QUESTIONS / COACH_MAX_TOKENS 这些默认值
    必须来自 .env；等 parse_args 之后才 load_env_file 就晚了（这是原先真实存在的失效点）。
    """
    for i, a in enumerate(argv):
        if a == "--env-file" and i + 1 < len(argv):
            return argv[i + 1], config.load_env_file(argv[i + 1])
        if a.startswith("--env-file="):
            path = a.split("=", 1)[1]
            return path, config.load_env_file(path)
    return None, 0


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv[1:]) if argv is None else list(argv)
    _env_path, _env_loaded = preload_env_file(argv)
    ap = argparse.ArgumentParser(description="多智能体求职陪练 Agent - 本地演示入口")
    ap.add_argument("--jd", required=True, help="JD 文本文件路径（.md/.txt/.docx）")
    ap.add_argument("--resume", required=True, help="简历文件路径（.docx/.md/.txt）")
    ap.add_argument("--answers", help="预置回答（.json 列表，或按空行分隔的 txt）")
    ap.add_argument("--interactive", action="store_true", help="真人练习模式（END 结束一题）")
    ap.add_argument("--questions", type=int, default=config.max_questions())
    ap.add_argument("--engine", choices=["local", "langgraph"], default="local")
    ap.add_argument("--run-id", default=None)
    ap.add_argument("--env-file", default=None, help="可选：读取 .env（不会写入仓库）")
    ap.add_argument("--seed", type=int, default=None, help="固定选题随机种子（同 seed 可复现同一套题）")
    ap.add_argument("--no-variety", action="store_true", help="关闭选题多样性（严格复现，评测/回归用）")
    ap.add_argument("--model", default=None, help="覆盖模型，逗号分隔可给降级链，如 deepseek-chat,deepseek-reasoner")
    ap.add_argument("--base-url", default=None, help="覆盖 OpenAI 兼容端点，如 https://api.deepseek.com")
    ap.add_argument("--eval", action="store_true", help="跑完自动执行评测 harness")
    ap.add_argument("--weight", default="", help='按 JD 定制要求权重，如 "R1=2,RAG=1.5,R3=0.5"')
    ap.add_argument("--html", action="store_true",
                    help="额外导出一份可打印 HTML（浏览器 Ctrl+P 直接存 PDF，零依赖）")
    ap.add_argument("--hitl", action="store_true",
                    help="人工确认闸门：弱点画像先落 pending，跑完后由你确认是否写入跨场次记忆")
    ap.add_argument("--reps", type=int, default=None,
                    help="每题模型采样次数（默认 1；评测/回归建议 3，取中位数以降低单次采样抖动）")
    args = ap.parse_args(argv)

    bad = check_inputs(args.jd, args.resume, args.answers)
    if bad:
        print_problems(bad)
        return 2

    if args.interactive and not sys.stdin.isatty():
        print("[提示] --interactive 需要在能打字的终端里跑（当前 stdin 不是交互终端），已改为按 --answers 作答。")
        args.interactive = False

    loaded = _env_loaded if (args.env_file and _env_path) else config.load_env_file(args.env_file)
    if args.model:                       # 显式参数优先于 .env
        os.environ["COACH_MODEL"] = args.model
    if args.base_url:
        os.environ["COACH_BASE_URL"] = args.base_url
    if args.reps:
        os.environ["COACH_SCORE_REPS"] = str(args.reps)
    if args.hitl:
        os.environ["COACH_HITL"] = "1"
    req_weights = {}
    for part in (args.weight or "").split(","):
        if "=" in part:
            k, _, v = part.partition("=")
            try:
                req_weights[k.strip()] = float(v)
            except ValueError:
                print(f"[warn] --weight 里的 {part!r} 不是合法倍数，已忽略")
    if req_weights:
        print("[评测参数] 要求权重: " + ", ".join(f"{k}x{v}" for k, v in req_weights.items()))
    config.ensure_dirs()
    jd_text = load_text(args.jd)
    resume_text = load_text(args.resume)
    print(f"[env] 载入变量 {loaded} 个 | 模式: {'离线(mock)' if config.offline() else '在线(真实模型)'}")
    print(f"[模型] {config.provider_summary()}")
    print(f"[评测参数] temperature={config.temperature()} | 注入防护 guard={os.getenv('COACH_GUARD', '2')}"
          f" | 每题采样 reps={config.score_reps()}"
          + ("  ← reps>1：分数取多次中位数，可比较性更好、成本×reps" if config.score_reps() > 1 else ""))

    answers, by_cat = load_answers(args.answers)
    if args.answers and not answers and not by_cat:
        print(f"[提示] 回答文件 {args.answers} 里没读到作答内容：本次会只出题、不评分（想评分请用 --interactive 现场作答）。")

    if args.engine == "langgraph":
        from coach.graph import run_with_langgraph
        preset = None
        if by_cat and not answers:      # 按类别回答需要先知道选题，物化成有序回答
            planner = CoachEngine(run_id=args.run_id, max_questions=args.questions, use_llm=False)
            preset = planner.plan(jd_text, resume_text).questions
            answers = [by_cat.get(q.qid) or by_cat.get(q.category, by_cat.get("default", ""))
                       for q in preset]
        report, run_id = run_with_langgraph(jd_text, resume_text, answers, run_id=args.run_id,
                                            max_questions=args.questions, preset_questions=preset)
    else:
        # 真人练习（--interactive）默认开启"多样 + 避开最近做过的题"；脚本演示保持确定性
        eng = CoachEngine(run_id=args.run_id, max_questions=args.questions, hitl=args.hitl,
                          req_weights=req_weights or None,
                          variety=(args.interactive and not args.no_variety) or args.seed is not None,
                          avoid_repeat=args.interactive, seed=args.seed)
        print(f"[model] {eng.model} | 题库 {len(eng.bank)} 条"
              + (f" | 选题: 多样(seed={args.seed})" if args.seed is not None
                 else (" | 选题: 多样+避开近 3 场" if (args.interactive and not args.no_variety)
                       else " | 选题: 确定性")))
        provider = None
        if args.interactive:
            provider = interactive_provider
        elif by_cat:
            # 按类别作答：追问（-F 结尾）不再作答，避免用同一段话重复打分
            def provider(q):
                base = q.qid.split("-F")[0]
                if q.qid != base:            # 追问题目：优先用专门的补充回答，没有就跳过
                    return by_cat.get(q.qid) or by_cat.get("__followup__", "")
                return by_cat.get(q.qid) or by_cat.get(q.category, by_cat.get("default", ""))
        report = eng.run(jd_text, resume_text, answers=answers, answer_provider=provider)
        run_id = eng.run_id

    print("\n" + "=" * 70)
    print(f"run_id        : {run_id}")
    print(f"岗位          : {report.jd_title}")
    print(f"要求覆盖率    : {report.match.score:.0%}" if report.match else "要求覆盖率    : -")
    print(f"综合得分      : {report.overall:.0%}")
    print(f"答题轮次      : {len(report.turns)}")
    print("逐题：")
    for t in report.turns:
        sc = t.scorecard
        if sc:
            print(f"  - {t.qid:9s} {sc.total:.0%} {sc.verdict:6s} 缺失: {'、'.join(sc.missing[:3]) or '-'}")
        else:
            print(f"  - {t.qid} 未评分")
    print(f"报告          : {report.report_path}")
    print(f"JSON          : {config.RESULTS_DIR}/run_{run_id}.json")
    print(f"trace         : {config.RUNS_DIR}/{run_id}/trace.jsonl")

    if args.html and getattr(report, "report_path", None):
        from coach.tools.export_html import export_html
        hp = export_html(report.report_path)
        print("\n可打印 HTML : %s（浏览器打开后 Ctrl+P 存 PDF）" % hp)
    if args.hitl and report.pending_review:
        print("\n" + "=" * 70)
        print("待你确认：以下弱点画像将写入【跨场次记忆】，会长期影响后续场次的练习重点")
        for i, w in enumerate(report.weaknesses, 1):
            print(f"  {i}. {w.tag} × {w.count} [{w.severity}]")
        try:
            raw = input("全部入库输 y ｜ 全部驳回输 n ｜ 只入库部分输编号（如 1,3）：").strip().lower()
        except EOFError:
            raw = "n"
        if raw == "y":
            approved = [w.tag for w in report.weaknesses]
        elif raw == "n":
            approved = []
        else:
            picks = {int(x) for x in re.findall(r"\d+", raw)}
            approved = [w.tag for i, w in enumerate(report.weaknesses, 1) if i in picks]
        if locals().get("eng") is not None:
            res = eng.resume_review(approved)
            print(f"[HITL] 已提交：入库 {res['committed_turns']} 轮；批准弱点 {len(res['approved'])} 条，"
                  f"驳回 {len(report.weaknesses) - len(res['approved'])} 条（驳回的不进画像）")
        else:
            print("[HITL] 这是 LangGraph 入口，请用 engine 入口跑 --hitl，或调用 /api/review/{}".format(run_id))
        print("=" * 70)

    try:
        from coach.eval import cost_report as _cost
        _agg = _cost.aggregate(run_id)
        _t = _agg["totals"]
        print(f"LLM 用量      : token {_t.get('total_tokens', 0)}"
              f"（prompt {_t.get('prompt_tokens', 0)} / completion {_t.get('completion_tokens', 0)}）")
        print(f"耗时          : LLM 等待 {_t.get('llm_ms', 0)/1000:.1f}s / 墙钟 {_t.get('wall_ms', 0)/1000:.1f}s")
        print(f"每次问答      : token {_agg['per_turn']['total_tokens']} | {_agg['per_turn']['ms']} ms")
        print(f"成本账        : {_cost.write(run_id)}")
    except Exception as e:
        print(f"[warn] 成本账生成失败：{type(e).__name__}: {e}")

    if args.eval:
        print("\n[eval] " + json.dumps(harness.evaluate(run_id), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
