"""生成"选题质量盲评"表：让关键词法与混合检索法在同一份 JD 上公平比一次。

背景：混合检索已实现，但两个**代理指标**结论相反（关键词重叠度支持关键词法、语义相似度支持混合法），
代理指标互相矛盾时不下结论 —— 需要人给"这道题贴不贴岗"打分。

做法（关键：**盲评**）：
  同一份 JD 分别用两种策略各选 top-k，取并集后**打乱顺序**写进表格，
  表里不标哪道题来自哪种策略（避免"知道来源就偏向"）。
  他只需要给每道题的"贴岗程度"打 1~3 分，我再用脚本按策略分组求均值比高低。

用法：
  python -m coach.eval.build_selector_sheet --k 8 --out results/selector_review.xlsx
  python -m coach.eval.review_selector_sheet results/selector_review.xlsx   # 填完后算
"""
from __future__ import annotations

import argparse
import json
import random
import sys
from datetime import datetime
from pathlib import Path

from .. import config
from ..tools import load_bank, parse_jd, plan_questions


def build(jd_text: str, k: int = 8, seed: int = 7) -> dict:
    bank = load_bank()
    prof = parse_jd(jd_text, llm=None)          # 规则解析：这一步不花 token
    kw = plan_questions(prof, k=k, bank=bank, variety=False)
    try:
        from ..retrieval import build_selector
        sel = build_selector(bank)
        hy = plan_questions(prof, k=k, bank=bank, selector=sel, variety=False)
        hy_err = ""
    except Exception as e:                      # 本机 bge 快照缺失时只出关键词一列，不装作有结论
        hy, hy_err = [], f"{type(e).__name__}: {e}"
    kw_ids = [q.qid for q in kw]
    hy_ids = [q.qid for q in hy]
    union = []
    for q in list(kw) + list(hy):
        if q.qid not in [x["qid"] for x in union]:
            union.append({"qid": q.qid, "category": q.category, "text": q.text,
                          "target_req_id": q.target_req_id,
                          "keywords": [kw2 for d in q.rubric.dims for kw2 in (d.keywords or [])][:6]})
    rng = random.Random(seed)
    rng.shuffle(union)                          # 打乱：不让"顺序"暗示策略
    return {"jd_title": prof.title, "k": k, "keyword_ids": kw_ids, "hybrid_ids": hy_ids,
            "hybrid_error": hy_err, "questions": union,
            "ts": datetime.now().strftime("%Y-%m-%d %H:%M")}


def write_xlsx(data: dict, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "选题盲评"
    head = ["序号", "题号", "类别", "题目", "参考要点（判贴岗用）", "贴岗程度 1-3（只填这列）", "备注"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    fill = PatternFill("solid", fgColor="FFF2CC")
    for i, q in enumerate(data["questions"], 1):
        ws.append([i, q["qid"], q["category"], q["text"], "、".join(q["keywords"]), None, ""])
        ws.cell(row=i + 1, column=6).fill = fill
    for col, w in zip("ABCDEFG", (6, 12, 12, 52, 46, 16, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("怎么填（口径）")
    for row in [
        ["这是干什么的", "混合检索实现了，但两个代理指标结论相反，需要人给'这道题贴不贴岗'打分才能决定默认策略"],
        ["只有一列需要填", "第 6 列『贴岗程度 1-3』（黄底）：3=很贴岗，2=沾边，1=不贴岗"],
        ["盲评", "表里题目顺序被打乱，且不标哪道题来自哪种策略 —— 避免知道来源就偏向"],
        ["不要留空", "没把握就按直觉填，留空该题不参与统计"],
        ["填完怎么办", "存盘后说一声，我用 review_selector_sheet 算两种策略各自的均值并给结论"],
        ["本表来源 JD", data["jd_title"]],
        ["生成时间", data["ts"]],
        ["混合检索加载", "成功" if not data["hybrid_error"] else ("失败：" + data["hybrid_error"][:80])],
    ]:
        ws2.append(row)
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 96
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="生成选题质量盲评表")
    ap.add_argument("--k", type=int, default=8, help="每种策略取前 k 道（并集后一起评）")
    ap.add_argument("--jd", default=None, help="JD 文本文件，默认用 data/jd_sample.md")
    ap.add_argument("--out", default=None)
    a = ap.parse_args(list(sys.argv[1:]) if argv is None else list(argv))
    jd = (Path(a.jd) if a.jd else Path(config.ROOT) / "data" / "jd_sample.md").read_text(encoding="utf-8")
    data = build(jd, k=a.k)
    out = Path(a.out) if a.out else Path(config.RESULTS_DIR) / "selector_review.xlsx"
    write_xlsx(data, out)
    meta = out.with_suffix(".meta.json")
    meta.write_text(json.dumps({"keyword_ids": data["keyword_ids"], "hybrid_ids": data["hybrid_ids"],
                                "hybrid_error": data["hybrid_error"], "k": a.k, "ts": data["ts"]},
                               ensure_ascii=False, indent=2), encoding="utf-8")
    print("关键词法 top%d: %s" % (a.k, data["keyword_ids"]))
    print("混合法   top%d: %s" % (a.k, data["hybrid_ids"]))
    if data["hybrid_error"]:
        print("⚠ 混合检索加载失败：%s" % data["hybrid_error"][:120])
    print("待评题目 %d 道（并集，已打乱）" % len(data["questions"]))
    print("表格: %s" % out)
    print("映射: %s（他填完后算均值要用）" % meta)
    return 0


if __name__ == "__main__":
    sys.exit(main())
