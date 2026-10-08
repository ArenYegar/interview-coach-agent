"""读回"选题盲评"表，按策略分组算均值，给出是否切换默认策略的结论。

用法：
    python -m coach.eval.review_selector_sheet results/selector_review.xlsx
    # 同目录需要 build_selector_sheet 生成的 selector_review.meta.json（记录每道题来自哪种策略）

判定口径（先定后跑）：
    两组各自算"贴岗程度"均值（1~3）；混合法比关键词法高出 **≥0.3 分**才建议切换默认策略，
    否则保持关键词法（现状）。差值不到 0.3 视为噪声 —— 别为了一个说不清的数改默认值。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def read_scores(xlsx: Path) -> dict:
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True).worksheets[0]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        qid, sc = str(row[1]), row[5]
        try:
            v = int(sc)
        except (TypeError, ValueError):
            continue
        if 1 <= v <= 3:
            out[qid] = v
    return out


def analyze(xlsx: Path, meta: Path | None = None, threshold: float = 0.3) -> dict:
    meta = meta or xlsx.with_suffix(".meta.json")
    if not meta.exists():
        raise SystemExit(f"找不到映射文件 {meta}（先跑 build_selector_sheet 生成）")
    m = json.loads(meta.read_text(encoding="utf-8"))
    scores = read_scores(xlsx)
    kw, hy = set(m["keyword_ids"]), set(m["hybrid_ids"])
    def mean(ids):
        vals = [scores[q] for q in ids if q in scores]
        return round(sum(vals) / len(vals), 3) if vals else None
    mk, mh = mean(kw), mean(hy)
    rated = len(scores)
    conclusion = "样本不足（评过的题 < 8 道）→ 不下结论"
    if mk is not None and mh is not None and rated >= 8:
        d = round(mh - mk, 3)
        conclusion = (f"混合法均值高 {d} 分（≥{threshold}）→ 建议把默认策略切到混合检索"
                      if d >= threshold else
                      f"差值 {d} 分（< {threshold}）→ 保持默认（关键词法），不为噪声改默认值")
    return {"rated": rated, "keyword_mean": mk, "hybrid_mean": mh,
            "keyword_n": len([q for q in kw if q in scores]), "hybrid_n": len([q for q in hy if q in scores]),
            "threshold": threshold, "conclusion": conclusion, "scores": scores}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="读回选题盲评表并给结论")
    ap.add_argument("xlsx")
    ap.add_argument("--threshold", type=float, default=0.3)
    a = ap.parse_args(list(sys.argv[1:]) if argv is None else list(argv))
    r = analyze(Path(a.xlsx), threshold=a.threshold)
    print("已评 %d 道题" % r["rated"])
    print("  关键词法: 均值 %s（%d 道）" % (r["keyword_mean"], r["keyword_n"]))
    print("  混合检索: 均值 %s（%d 道）" % (r["hybrid_mean"], r["hybrid_n"]))
    print("  结论: %s" % r["conclusion"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
